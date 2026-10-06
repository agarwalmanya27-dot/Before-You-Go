import mysql.connector
from openpyxl import load_workbook

# MySQL connection
db = mysql.connector.connect(
    host="localhost",
    user="root",
    password="kswork07",
    database="byg"
)

cursor = db.cursor()

# Excel file read
workbook = load_workbook("destinations.xlsx")
sheet = workbook.active

# Excel ki har row ko MySQL mein insert karna
for row in sheet.iter_rows(min_row=2, values_only=True):
    dest_id, dest_name, state, vibe, descrip, img_path = row

    sql = """
    INSERT INTO destinations
    (dest_id, dest_name, state, vibe, descrip, img_path)
    VALUES (%s, %s, %s, %s, %s, %s)
    """

    cursor.execute(sql, (
        dest_id,
        dest_name,
        state,
        vibe,
        descrip,
        img_path
    ))

db.commit()

print("Destinations data successfully imported!")

cursor.close()
db.close()